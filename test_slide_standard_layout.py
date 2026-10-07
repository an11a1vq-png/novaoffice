import io
import unittest
from fastapi.testclient import TestClient
import pptx
from pptx.util import Inches

from app.main import app
from app.services.pptx_service import PptxService

class TestSlideStandardLayout(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app)

    def test_slide_html_has_standard_layout_option(self):
        res = self.client.get("/slide")
        self.assertEqual(res.status_code, 200)
        self.assertIn('<option value="standard">', res.text)
        self.assertIn('Tiêu chuẩn (Tiêu đề & Nội dung)', res.text)

    def test_slide_js_has_standard_layout_handlers(self):
        res = self.client.get("/static/js/slide.js")
        self.assertEqual(res.status_code, 200)
        js_code = res.text
        self.assertIn("case 'standard':", js_code)
        self.assertIn("updateSlideData('content', this.innerText)", js_code)
        self.assertIn("whitespace-pre-wrap", js_code)

    def test_pptx_import_and_export_with_notes(self):
        prs = pptx.Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        
        # Add Title textbox
        tb_title = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
        tb_title.text_frame.text = "Ví dụ hoàn chỉnh về Cookie"

        # Add Body textbox
        tb_body = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(10), Inches(4))
        tb_body.text_frame.text = "<?php\nsetcookie('user', 'John', time() + 3600);\n?>"

        # Add speaker notes
        slide.notes_slide.notes_text_frame.text = "Lưu ý cấu hình cookie an toàn với HttpOnly."

        # Save to buffer
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)

        # Import
        imported_deck = PptxService.import_from_pptx(buf)
        self.assertIn("slides", imported_deck)
        self.assertEqual(len(imported_deck["slides"]), 1)
        
        first_slide = imported_deck["slides"][0]
        self.assertEqual(first_slide["title"], "Ví dụ hoàn chỉnh về Cookie")
        self.assertIn("setcookie", first_slide["content"])
        self.assertEqual(first_slide["layout"], "standard")
        self.assertEqual(first_slide.get("notes"), "Lưu ý cấu hình cookie an toàn với HttpOnly.")

        # Re-export to PPTX
        export_buf = PptxService.export_to_pptx(imported_deck, "Test_Cookie")
        self.assertGreater(export_buf.getbuffer().nbytes, 1000)

        # Re-import to verify round-trip
        reimported = PptxService.import_from_pptx(export_buf)
        self.assertEqual(len(reimported["slides"]), 1)
        self.assertEqual(reimported["slides"][0]["title"], "Ví dụ hoàn chỉnh về Cookie")
        self.assertIn("setcookie", reimported["slides"][0]["content"])
        self.assertEqual(reimported["slides"][0].get("notes"), "Lưu ý cấu hình cookie an toàn với HttpOnly.")

if __name__ == "__main__":
    unittest.main()
