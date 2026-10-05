import io
from pathlib import Path
from bs4 import BeautifulSoup
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.text.paragraph import Paragraph
from docx.table import Table
import html

try:
    import mammoth
    MAMMOTH_AVAILABLE = True
except ImportError:
    MAMMOTH_AVAILABLE = False

class DocxService:
    @staticmethod
    def html_to_docx(html_content: str, title: str = "Tài liệu NovaDoc") -> io.BytesIO:
        doc = docx.Document()
        
        # Set margins (1 inch)
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        soup = BeautifulSoup(html_content, "html.parser")
        
        # Traverse elements
        for element in soup.children:
            if not hasattr(element, "name") or element.name is None:
                text = str(element).strip()
                if text:
                    doc.add_paragraph(text)
                continue

            tag_name = element.name.lower()

            if tag_name in ["h1", "h2", "h3", "h4", "h5", "h6"]:
                level = int(tag_name[1])
                p = doc.add_heading(level=level)
                DocxService._apply_inline_formatting(element, p)

            elif tag_name == "p":
                p = doc.add_paragraph()
                align = element.get("style", "")
                if "text-align: center" in align:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif "text-align: right" in align:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                elif "text-align: justify" in align:
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                DocxService._apply_inline_formatting(element, p)

            elif tag_name in ["ul", "ol"]:
                is_ordered = (tag_name == "ol")
                for li in element.find_all("li", recursive=False):
                    style = "List Number" if is_ordered else "List Bullet"
                    p = doc.add_paragraph(style=style)
                    DocxService._apply_inline_formatting(li, p)

            elif tag_name == "blockquote":
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.5)
                run = p.add_run(element.get_text())
                run.italic = True
                run.font.color.rgb = RGBColor(100, 100, 100)

            elif tag_name == "table":
                rows = element.find_all("tr")
                if rows:
                    num_rows = len(rows)
                    num_cols = max(len(r.find_all(["td", "th"])) for r in rows)
                    table = doc.add_table(rows=num_rows, cols=num_cols)
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    table.style = 'Table Grid'
                    
                    for row_idx, r in enumerate(rows):
                        cells = r.find_all(["td", "th"])
                        for col_idx, c in enumerate(cells):
                            if col_idx < num_cols:
                                cell = table.cell(row_idx, col_idx)
                                cell.text = c.get_text().strip()
                                if c.name == "th":
                                    for paragraph in cell.paragraphs:
                                        for run in paragraph.runs:
                                            run.bold = True
                doc.add_paragraph()  # spacing

            elif tag_name == "hr":
                p = doc.add_paragraph("──────────────────────────────────────────")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            else:
                p = doc.add_paragraph()
                DocxService._apply_inline_formatting(element, p)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer

    @staticmethod
    def _apply_inline_formatting(parent_tag, paragraph):
        for child in parent_tag.children:
            if isinstance(child, str):
                if child.strip() or child == " ":
                    paragraph.add_run(child)
            elif child.name in ["b", "strong"]:
                run = paragraph.add_run(child.get_text())
                run.bold = True
            elif child.name in ["i", "em"]:
                run = paragraph.add_run(child.get_text())
                run.italic = True
            elif child.name == "u":
                run = paragraph.add_run(child.get_text())
                run.underline = True
            elif child.name in ["s", "strike"]:
                run = paragraph.add_run(child.get_text())
                run.font.strike = True
            elif child.name == "span":
                run = paragraph.add_run(child.get_text())
            elif child.name == "a":
                run = paragraph.add_run(child.get_text())
                run.underline = True
                run.font.color.rgb = RGBColor(0, 102, 204)
            else:
                paragraph.add_run(child.get_text())

    @staticmethod
    def docx_to_html(file_stream: io.BytesIO) -> str:
        """
        Converts a Word .docx stream into rich, semantic HTML preserving formatting,
        Vietnamese characters, tables in original sequential order, bold, italics, and lists.
        """
        if MAMMOTH_AVAILABLE:
            try:
                # Seek to start
                file_stream.seek(0)
                result = mammoth.convert_to_html(file_stream)
                raw_html = result.value

                # Post-process with BeautifulSoup to apply modern table styles and clean layout
                soup = BeautifulSoup(raw_html, "html.parser")
                
                # Format tables nicely
                for table in soup.find_all("table"):
                    table["border"] = "1"
                    table["style"] = "width: 100%; border-collapse: collapse; margin: 18px 0; font-size: 13.5px;"
                    for th in table.find_all("th"):
                        th["style"] = "border: 1px solid #cbd5e1; padding: 8px 10px; background-color: #f8fafc; font-weight: 600; text-align: left; vertical-align: top;"
                    for td in table.find_all("td"):
                        td["style"] = "border: 1px solid #cbd5e1; padding: 8px 10px; vertical-align: top;"
                
                # Ensure images are constrained to page width
                for img in soup.find_all("img"):
                    img["style"] = "max-width: 100%; height: auto; border-radius: 4px; margin: 10px 0;"

                # Format headings with clean margins
                for h in soup.find_all(["h1", "h2", "h3", "h4"]):
                    h["style"] = "margin: 16px 0 8px 0; font-weight: bold; color: #1e293b;"

                # Format lists
                for ul in soup.find_all("ul"):
                    ul["style"] = "margin: 10px 0 14px 20px; list-style-type: disc;"
                for ol in soup.find_all("ol"):
                    ol["style"] = "margin: 10px 0 14px 20px; list-style-type: decimal;"

                return str(soup)
            except Exception as e:
                # Fallback to python-docx sequential walker
                pass

        # Robust Fallback using python-docx sequential document walker
        file_stream.seek(0)
        doc = docx.Document(file_stream)
        html_parts = []

        for child in doc.element.body:
            if child.tag.endswith('p'):
                p = Paragraph(child, doc)
                text = p.text.strip()
                if not text:
                    continue

                # Build paragraph with run formatting
                run_html = []
                for run in p.runs:
                    r_text = html.escape(run.text)
                    if not r_text:
                        continue
                    if run.bold:
                        r_text = f"<strong>{r_text}</strong>"
                    if run.italic:
                        r_text = f"<em>{r_text}</em>"
                    if run.underline:
                        r_text = f"<u>{r_text}</u>"
                    run_html.append(r_text)

                content = "".join(run_html) if run_html else html.escape(text)

                if p.style.name.startswith("Heading 1"):
                    html_parts.append(f"<h1>{content}</h1>")
                elif p.style.name.startswith("Heading 2"):
                    html_parts.append(f"<h2>{content}</h2>")
                elif p.style.name.startswith("Heading 3"):
                    html_parts.append(f"<h3>{content}</h3>")
                elif p.style.name.startswith("List"):
                    html_parts.append(f"<li>{content}</li>")
                else:
                    html_parts.append(f"<p>{content}</p>")

            elif child.tag.endswith('tbl'):
                t = Table(child, doc)
                table_html = ["<table border='1' style='width: 100%; border-collapse: collapse; margin: 18px 0;'>"]
                for row_idx, row in enumerate(t.rows):
                    table_html.append("<tr>")
                    for cell in row.cells:
                        tag = "th" if row_idx == 0 else "td"
                        style = "border: 1px solid #cbd5e1; padding: 8px 10px; background-color: #f8fafc; font-weight: bold;" if row_idx == 0 else "border: 1px solid #cbd5e1; padding: 8px 10px;"
                        table_html.append(f"<{tag} style='{style}'>{html.escape(cell.text)}</{tag}>")
                    table_html.append("</tr>")
                table_html.append("</table>")
                html_parts.append("".join(table_html))

        return "\n".join(html_parts)
