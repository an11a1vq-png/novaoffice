import io
import re
from typing import Any, Dict, List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class XlsxService:
    @staticmethod
    def export_to_xlsx(sheet_data_dict: Dict[str, Any], title: str = "Bang_Tinh_NovaOffice") -> io.BytesIO:
        wb = openpyxl.Workbook()
        # remove default sheet
        wb.remove(wb.active)

        sheets = sheet_data_dict.get("sheets", {})
        if not sheets:
            sheets = {"Sheet1": {"data": {}}}

        thin_border = Border(
            left=Side(style='thin', color='D0D5DD'),
            right=Side(style='thin', color='D0D5DD'),
            top=Side(style='thin', color='D0D5DD'),
            bottom=Side(style='thin', color='D0D5DD')
        )

        for sheet_name, sheet_content in sheets.items():
            ws = wb.create_sheet(title=str(sheet_name)[:31])
            data = sheet_content.get("data", {})

            for coord, cell_info in data.items():
                if not re.match(r"^[A-Z]+[0-9]+$", coord):
                    continue
                try:
                    cell = ws[coord]
                    # check if value is formula
                    val = cell_info.get("formula") or cell_info.get("value")
                    cell.value = val

                    # styles
                    is_bold = bool(cell_info.get("bold", False))
                    is_italic = bool(cell_info.get("italic", False))
                    text_color = cell_info.get("color")
                    bg_color = cell_info.get("bg")
                    align_val = cell_info.get("align", "left")

                    font_kwargs = {"bold": is_bold, "italic": is_italic, "name": "Calibri", "size": 11}
                    if text_color and text_color.startswith("#") and len(text_color) == 7:
                        font_kwargs["color"] = text_color[1:]
                    cell.font = Font(**font_kwargs)

                    if bg_color and bg_color.startswith("#") and len(bg_color) == 7:
                        cell.fill = PatternFill(start_color=bg_color[1:], end_color=bg_color[1:], fill_type="solid")

                    cell.alignment = Alignment(horizontal=align_val, vertical="center")
                    cell.border = thin_border
                except Exception:
                    continue

            # adjust column widths slightly
            for col in ws.columns:
                max_len = 10
                col_letter = get_column_letter(col[0].column)
                for c in col:
                    if c.value is not None:
                        max_len = max(max_len, min(len(str(c.value)) + 3, 30))
                ws.column_dimensions[col_letter].width = max_len

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    @staticmethod
    def import_from_xlsx(file_stream: io.BytesIO) -> Dict[str, Any]:
        wb = openpyxl.load_workbook(file_stream, data_only=False)
        result_sheets = {}

        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            sheet_cells = {}
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        val = cell.value
                        cell_dict: Dict[str, Any] = {}
                        if str(val).startswith("="):
                            cell_dict["formula"] = str(val)
                        else:
                            cell_dict["value"] = val

                        if cell.font:
                            if cell.font.bold:
                                cell_dict["bold"] = True
                            if cell.font.italic:
                                cell_dict["italic"] = True
                            if cell.font.color and cell.font.color.rgb and isinstance(cell.font.color.rgb, str):
                                cell_dict["color"] = f"#{cell.font.color.rgb[-6:]}"

                        if cell.fill and cell.fill.start_color and cell.fill.start_color.rgb and isinstance(cell.fill.start_color.rgb, str):
                            cell_dict["bg"] = f"#{cell.fill.start_color.rgb[-6:]}"

                        sheet_cells[cell.coordinate] = cell_dict

            result_sheets[sheetname] = {"data": sheet_cells}

        return {
            "activeSheet": wb.sheetnames[0] if wb.sheetnames else "Trang tính 1",
            "sheets": result_sheets
        }
