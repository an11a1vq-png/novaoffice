import io
from typing import Any, Dict, List, Optional
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def rgb_to_hex(rgb) -> Optional[str]:
    try:
        return f"#{rgb[0]:02x}{rgb[1]:02x}{rgb[2]:02x}".upper()
    except Exception:
        return None

def hex_to_rgb(hex_str: Optional[str]) -> Optional[RGBColor]:
    if not hex_str:
        return None
    cleaned = hex_str.lstrip("#")
    if len(cleaned) == 6:
        try:
            return RGBColor(int(cleaned[0:2], 16), int(cleaned[2:4], 16), int(cleaned[4:6], 16))
        except Exception:
            return None
    return None

THEME_COLORS = {
    "modern-dark": {
        "bg": RGBColor(15, 23, 42),       # slate 900
        "title": RGBColor(248, 250, 252),  # slate 50
        "body": RGBColor(203, 213, 225),   # slate 300
        "accent": RGBColor(56, 189, 248)   # sky 400
    },
    "clean-light": {
        "bg": RGBColor(255, 255, 255),
        "title": RGBColor(30, 41, 59),
        "body": RGBColor(71, 85, 105),
        "accent": RGBColor(37, 99, 235)
    },
    "corporate-blue": {
        "bg": RGBColor(240, 249, 255),
        "title": RGBColor(12, 74, 110),
        "body": RGBColor(30, 58, 138),
        "accent": RGBColor(2, 132, 199)
    },
    "emerald": {
        "bg": RGBColor(240, 253, 244),
        "title": RGBColor(6, 78, 59),
        "body": RGBColor(22, 101, 52),
        "accent": RGBColor(16, 185, 129)
    }
}

class PptxService:
    @staticmethod
    def export_to_pptx(deck_data: Dict[str, Any], default_title: str = "Bai_Thuyet_Trinh") -> io.BytesIO:
        prs = pptx.Presentation()
        prs.slide_width = Inches(13.333)  # 16:9 widescreen
        prs.slide_height = Inches(7.5)

        blank_slide_layout = prs.slide_layouts[6]
        theme_name = deck_data.get("theme", "modern-dark")
        palette = THEME_COLORS.get(theme_name, THEME_COLORS["modern-dark"])

        slides_list = deck_data.get("slides", [])
        if not slides_list:
            slides_list = [{
                "title": default_title,
                "subtitle": "Tạo bằng NovaSlide",
                "layout": "title-slide"
            }]

        for slide_info in slides_list:
            slide = prs.slides.add_slide(blank_slide_layout)
            
            # Fill background
            background = slide.background
            fill = background.fill
            fill.solid()
            fill.fore_color.rgb = palette["bg"]

            layout = slide_info.get("layout", "standard")
            title_text = slide_info.get("title", "")
            subtitle_text = slide_info.get("subtitle", "")
            content_text = slide_info.get("content", "")
            bullets = slide_info.get("bullets", [])

            if layout == "shapes" and slide_info.get("shapes"):
                # Background
                bg_hex = slide_info.get("bg_color")
                bg_rgb = hex_to_rgb(bg_hex) or palette["bg"]
                fill = slide.background.fill
                fill.solid()
                fill.fore_color.rgb = bg_rgb

                for sh_data in slide_info.get("shapes", []):
                    s_type = sh_data.get("type", "card")
                    left_in = Inches(prs.slide_width.inches * (sh_data.get("left", 0) / 100))
                    top_in = Inches(prs.slide_height.inches * (sh_data.get("top", 0) / 100))
                    width_in = Inches(prs.slide_width.inches * (sh_data.get("width", 0) / 100))
                    height_in = Inches(prs.slide_height.inches * (sh_data.get("height", 0) / 100))

                    if s_type == "line":
                        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_in, top_in, width_in, max(height_in, Inches(0.04)))
                        shp.line.fill.background()
                        fill_c = hex_to_rgb(sh_data.get("fill") or sh_data.get("border")) or palette["accent"]
                        shp.fill.solid()
                        shp.fill.fore_color.rgb = fill_c

                    elif s_type == "table" and sh_data.get("table_data"):
                        tbl_data = sh_data.get("table_data", [])
                        rows_cnt = len(tbl_data)
                        cols_cnt = len(tbl_data[0]) if rows_cnt > 0 else 0
                        if rows_cnt > 0 and cols_cnt > 0:
                            tbl_shape = slide.shapes.add_table(rows_cnt, cols_cnt, left_in, top_in, width_in, height_in)
                            tbl = tbl_shape.table
                            for r_idx, row_list in enumerate(tbl_data):
                                for c_idx, cell_val in enumerate(row_list):
                                    if c_idx < cols_cnt:
                                        tbl.cell(r_idx, c_idx).text = str(cell_val)

                    else:
                        fill_rgb = hex_to_rgb(sh_data.get("fill"))
                        border_rgb = hex_to_rgb(sh_data.get("border"))

                        if fill_rgb or border_rgb:
                            shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_in, top_in, width_in, height_in)
                            if fill_rgb:
                                shp.fill.solid()
                                shp.fill.fore_color.rgb = fill_rgb
                            else:
                                shp.fill.background()

                            if border_rgb:
                                shp.line.color.rgb = border_rgb
                                shp.line.width = Pt(1.5)
                            else:
                                shp.line.fill.background()

                            tf = shp.text_frame
                        else:
                            tb = slide.shapes.add_textbox(left_in, top_in, width_in, height_in)
                            tf = tb.text_frame

                        tf.word_wrap = True
                        paragraphs = sh_data.get("paragraphs", [])
                        if paragraphs:
                            for p_idx, p_info in enumerate(paragraphs):
                                p = tf.paragraphs[0] if p_idx == 0 else tf.add_paragraph()
                                p.text = p_info.get("text", "")
                                if p_info.get("size"):
                                    p.font.size = Pt(p_info.get("size"))
                                if p_info.get("bold"):
                                    p.font.bold = True
                                if p_info.get("font"):
                                    p.font.name = p_info.get("font")
                                if p_info.get("color"):
                                    c_rgb = hex_to_rgb(p_info.get("color"))
                                    if c_rgb:
                                        p.font.color.rgb = c_rgb
                        elif sh_data.get("text"):
                            p = tf.paragraphs[0]
                            p.text = sh_data.get("text", "")
                            if sh_data.get("font_size"):
                                p.font.size = Pt(sh_data.get("font_size"))
                            if sh_data.get("font_color"):
                                c_rgb = hex_to_rgb(sh_data.get("font_color"))
                                if c_rgb:
                                    p.font.color.rgb = c_rgb

            elif layout == "title-slide":
                # Center Title & Subtitle
                tx_box = slide.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.33), Inches(3.2))
                tf = tx_box.text_frame
                tf.word_wrap = True

                p1 = tf.paragraphs[0]
                p1.text = title_text
                p1.alignment = PP_ALIGN.CENTER
                p1.font.size = Pt(48)
                p1.font.bold = True
                p1.font.color.rgb = palette["title"]

                if subtitle_text:
                    p2 = tf.add_paragraph()
                    p2.text = subtitle_text
                    p2.alignment = PP_ALIGN.CENTER
                    p2.font.size = Pt(24)
                    p2.font.color.rgb = palette["accent"]
                    p2.space_before = Pt(20)

            elif layout == "bullet-list" or bullets:
                # Title on top
                tx_box_title = slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.33), Inches(1.2))
                tf_title = tx_box_title.text_frame
                tf_title.word_wrap = True
                p_title = tf_title.paragraphs[0]
                p_title.text = title_text
                p_title.font.size = Pt(36)
                p_title.font.bold = True
                p_title.font.color.rgb = palette["title"]

                # Bullets below
                tx_box_body = slide.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(10.5), Inches(4.5))
                tf_body = tx_box_body.text_frame
                tf_body.word_wrap = True

                first = True
                for b_item in bullets:
                    p = tf_body.paragraphs[0] if first else tf_body.add_paragraph()
                    first = False
                    p.text = f"•  {b_item}"
                    p.font.size = Pt(24)
                    p.font.color.rgb = palette["body"]
                    p.space_after = Pt(16)

            elif layout == "process-3step":
                # Title on top
                tx_box_title = slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.33), Inches(1.2))
                tf_title = tx_box_title.text_frame
                tf_title.word_wrap = True
                p_title = tf_title.paragraphs[0]
                p_title.text = title_text or "Quy Trình 3 Bước"
                p_title.font.size = Pt(36)
                p_title.font.bold = True
                p_title.font.color.rgb = palette["title"]

                steps = [
                    (slide_info.get("col1Title", "Bước 1"), slide_info.get("col1Content", "")),
                    (slide_info.get("col2Title", "Bước 2"), slide_info.get("col2Content", "")),
                    (slide_info.get("col3Title", "Bước 3"), slide_info.get("col3Content", "")),
                ]
                card_w = Inches(3.5)
                card_h = Inches(4.5)
                card_y = Inches(2.2)
                for idx, (st_title, st_desc) in enumerate(steps):
                    card_x = Inches(1.0 + idx * 3.8)
                    tx_box = slide.shapes.add_textbox(card_x, card_y, card_w, card_h)
                    tf = tx_box.text_frame
                    tf.word_wrap = True
                    p_num = tf.paragraphs[0]
                    p_num.text = f"BƯỚC {idx + 1}"
                    p_num.font.size = Pt(14)
                    p_num.font.bold = True
                    p_num.font.color.rgb = palette["accent"]

                    p_t = tf.add_paragraph()
                    p_t.text = st_title
                    p_t.font.size = Pt(20)
                    p_t.font.bold = True
                    p_t.font.color.rgb = palette["title"]
                    p_t.space_before = Pt(8)

                    p_c = tf.add_paragraph()
                    p_c.text = st_desc
                    p_c.font.size = Pt(16)
                    p_c.font.color.rgb = palette["body"]
                    p_c.space_before = Pt(12)

            else:
                # Standard Title + Content
                tx_box_title = slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.33), Inches(1.2))
                tf_title = tx_box_title.text_frame
                tf_title.word_wrap = True
                p_title = tf_title.paragraphs[0]
                p_title.text = title_text
                p_title.font.size = Pt(36)
                p_title.font.bold = True
                p_title.font.color.rgb = palette["title"]

                tx_box_body = slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.33), Inches(4.5))
                tf_body = tx_box_body.text_frame
                tf_body.word_wrap = True
                p_body = tf_body.paragraphs[0]
                p_body.text = content_text or subtitle_text
                p_body.font.size = Pt(22)
                p_body.font.color.rgb = palette["body"]

            # Save speaker notes if present
            notes_text = slide_info.get("notes", "")
            if notes_text:
                try:
                    slide.notes_slide.notes_text_frame.text = str(notes_text)
                except Exception:
                    pass

        buffer = io.BytesIO()
        prs.save(buffer)
        buffer.seek(0)
        return buffer

    @staticmethod
    def import_from_pptx(stream: io.BytesIO) -> Dict[str, Any]:
        prs = pptx.Presentation(stream)
        slides_list = []
        sw = prs.slide_width
        sh_h = prs.slide_height

        for slide in prs.slides:
            # 1. Slide Background
            bg_hex = None
            try:
                if slide.background.fill.type and slide.background.fill.fore_color.rgb:
                    bg_hex = rgb_to_hex(slide.background.fill.fore_color.rgb)
            except Exception:
                pass

            title = ""
            texts = []
            shapes_list = []
            has_rich_shapes = False

            # Check title shape first
            title_shape = None
            try:
                if slide.shapes.title and slide.shapes.title.has_text_frame:
                    t = slide.shapes.title.text_frame.text.replace('\x0b', '\n').strip()
                    if t:
                        title = t
                        title_shape = slide.shapes.title
            except Exception:
                title_shape = None

            # Sort shapes in reading order (top then left)
            sorted_shapes = sorted(slide.shapes, key=lambda s: (getattr(s, "top", 0) or 0, getattr(s, "left", 0) or 0))

            for sh_idx, shape in enumerate(sorted_shapes):
                left_pct = round(shape.left / sw * 100, 2)
                top_pct = round(shape.top / sh_h * 100, 2)
                w_pct = round(shape.width / sw * 100, 2)
                h_pct = round(shape.height / sh_h * 100, 2)

                fill_hex = None
                try:
                    if shape.fill.type and shape.fill.fore_color and shape.fill.fore_color.rgb:
                        fill_hex = rgb_to_hex(shape.fill.fore_color.rgb)
                except Exception:
                    pass

                border_hex = None
                try:
                    if shape.line and shape.line.color and shape.line.color.rgb:
                        border_hex = rgb_to_hex(shape.line.color.rgb)
                except Exception:
                    pass

                if fill_hex or border_hex:
                    has_rich_shapes = True

                is_line = (h_pct <= 0.8 or w_pct <= 0.8) and (not shape.has_text_frame or not shape.text_frame.text.strip())
                is_code = False
                text = ""
                paras = []
                table_data = None

                if getattr(shape, "has_table", False):
                    table_data = []
                    for row in shape.table.rows:
                        table_data.append([c.text.replace('\x0b', '\n').strip() for c in row.cells])
                    shape_type = "table"
                    has_rich_shapes = True
                    for row in table_data:
                        texts.append(" | ".join(row))
                elif is_line:
                    shape_type = "line"
                    has_rich_shapes = True
                else:
                    if shape.has_text_frame:
                        text = shape.text_frame.text.replace('\x0b', '\n').strip()
                        if not title and text:
                            title = text
                        elif text:
                            texts.append(text)

                        for p_obj in shape.text_frame.paragraphs:
                            p_txt = p_obj.text.replace('\x0b', '\n').strip()
                            if not p_txt:
                                continue
                            p_col = None
                            p_bold = False
                            p_font = None
                            p_size = None
                            for r in p_obj.runs:
                                if not p_font and r.font.name:
                                    p_font = r.font.name
                                if not p_size and r.font.size:
                                    p_size = round(r.font.size.pt, 1)
                                if not p_col:
                                    try:
                                        if r.font.color and r.font.color.rgb:
                                            p_col = rgb_to_hex(r.font.color.rgb)
                                    except Exception:
                                        pass
                                if r.font.bold:
                                    p_bold = True
                            if p_font in ['Consolas', 'Courier New']:
                                is_code = True
                            paras.append({
                                "text": p_txt,
                                "color": p_col,
                                "bold": p_bold,
                                "font": p_font,
                                "size": p_size
                            })
                        if text.startswith(('//', 'using ', 'var ', 'namespace ', 'public class', '<?php')):
                            is_code = True

                    shape_type = "code" if is_code else "card"

                shapes_list.append({
                    "id": f"sh-{sh_idx}",
                    "type": shape_type,
                    "left": left_pct,
                    "top": top_pct,
                    "width": w_pct,
                    "height": h_pct,
                    "fill": fill_hex,
                    "border": border_hex,
                    "is_code": is_code,
                    "text": text,
                    "paragraphs": paras,
                    "table_data": table_data
                })

            content = "\n\n".join(texts)

            # Extract speaker notes if available
            notes = ""
            try:
                if getattr(slide, "has_notes_slide", False):
                    notes_frame = getattr(slide.notes_slide, "notes_text_frame", None)
                    if notes_frame and notes_frame.text:
                        notes = notes_frame.text.replace('\x0b', '\n').strip()
            except Exception:
                pass

            use_shapes = has_rich_shapes or len(shapes_list) > 2

            slide_dict = {
                "title": title or "Trang chiếu",
                "content": content,
                "layout": "shapes" if use_shapes else "standard",
                "bg_color": bg_hex or "#0F172A",
                "shapes": shapes_list
            }
            if notes:
                slide_dict["notes"] = notes

            slides_list.append(slide_dict)

        if not slides_list:
            slides_list = [{"title": "Trang chiếu mới", "content": "", "layout": "standard"}]
        return {"slides": slides_list, "theme": "modern-dark"}
