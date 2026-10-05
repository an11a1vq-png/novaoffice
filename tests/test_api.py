import unittest
from fastapi.testclient import TestClient
from app.main import app
from app.services.storage import StorageService

class TestNovaOfficeSuite(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    def test_hub_page_loads(self):
        res = self.client.get("/")
        self.assertEqual(res.status_code, 200)
        self.assertIn("NovaOffice", res.text)

    def test_app_pages_load(self):
        for route in ["/doc", "/sheet", "/slide", "/pdf"]:
            res = self.client.get(route)
            self.assertEqual(res.status_code, 200)

    def test_crud_document_lifecycle(self):
        # 1. Create Doc
        create_payload = {
            "title": "Tài liệu kiểm thử tự động",
            "type": "doc",
            "tags": ["test", "automated"]
        }
        res = self.client.post("/api/documents", json=create_payload)
        self.assertEqual(res.status_code, 200)
        data = res.json()
        doc_id = data["id"]
        self.assertEqual(data["title"], "Tài liệu kiểm thử tự động")
        self.assertEqual(data["type"], "doc")

        # 2. Retrieve Doc
        get_res = self.client.get(f"/api/documents/{doc_id}")
        self.assertEqual(get_res.status_code, 200)
        self.assertEqual(get_res.json()["id"], doc_id)

        # 3. Update Doc
        update_payload = {
            "title": "Tài liệu đã cập nhật",
            "content": {"html": "<p>Nội dung mới đã sửa</p>"}
        }
        update_res = self.client.put(f"/api/documents/{doc_id}", json=update_payload)
        self.assertEqual(update_res.status_code, 200)
        self.assertEqual(update_res.json()["title"], "Tài liệu đã cập nhật")

        # 4. Duplicate Doc
        dup_res = self.client.post(f"/api/documents/{doc_id}/duplicate")
        self.assertEqual(dup_res.status_code, 200)
        dup_id = dup_res.json()["id"]
        self.assertNotEqual(dup_id, doc_id)
        self.assertIn("Bản sao", dup_res.json()["title"])

        # 5. Delete both
        del1 = self.client.delete(f"/api/documents/{doc_id}")
        del2 = self.client.delete(f"/api/documents/{dup_id}")
        self.assertEqual(del1.status_code, 200)
        self.assertEqual(del2.status_code, 200)

    def test_export_docx(self):
        payload = {
            "title": "Bao_Cao_Test",
            "html_content": "<h1>Tiêu đề</h1><p>Đoạn văn kiểm thử</p>"
        }
        res = self.client.post("/api/export/docx", json=payload)
        self.assertEqual(res.status_code, 200)
        self.assertTrue(len(res.content) > 1000)

    def test_export_xlsx(self):
        payload = {
            "title": "Bang_Tinh_Test",
            "sheets": {
                "Sheet1": [
                    ["Tên", "Số lượng", "Giá"],
                    ["Sản phẩm A", 10, 50000],
                    ["Sản phẩm B", 5, 20000]
                ]
            }
        }
        res = self.client.post("/api/export/xlsx", json=payload)
        self.assertEqual(res.status_code, 200)
        self.assertTrue(len(res.content) > 1000)

    def test_export_pptx(self):
        payload = {
            "title": "Thuyet_Trinh_Test",
            "slides": [
                {
                    "title": "Slide 1",
                    "subtitle": "Phụ đề 1",
                    "layout": "title-slide"
                },
                {
                    "title": "Slide 2",
                    "layout": "bullet-list",
                    "bullets": ["Ý 1", "Ý 2"]
                }
            ]
        }
        res = self.client.post("/api/export/pptx", json=payload)
        self.assertEqual(res.status_code, 200)
        self.assertTrue(len(res.content) > 1000)

    def test_export_pdf(self):
        payload = {
            "title": "PDF_Test",
            "html_content": "<h1>Tiêu đề PDF</h1><p>Nội dung trang</p>"
        }
        res = self.client.post("/api/export/pdf", json=payload)
        self.assertEqual(res.status_code, 200)
        self.assertTrue(len(res.content) > 500)

    def test_docx_import_preserves_formatting(self):
        import io
        import docx
        from app.services.docx_service import DocxService

        doc = docx.Document()
        p = doc.add_paragraph()
        run = p.add_run("Văn bản in đậm tiếng Việt")
        run.bold = True
        
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Ô 1"
        table.cell(0, 1).text = "Ô 2"
        table.cell(1, 0).text = "Ô 3"
        table.cell(1, 1).text = "Ô 4"

        doc.add_paragraph("Đoạn văn kết thúc")

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)

        html_out = DocxService.docx_to_html(buf)
        self.assertIn("Văn bản in đậm tiếng Việt", html_out)
        self.assertTrue("strong" in html_out or "b" in html_out)
        self.assertIn("<table", html_out)
        self.assertIn("Ô 1", html_out)

    def test_ai_status_endpoint(self):
        res = self.client.get("/api/ai/status")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertIn("mode", data)
        self.assertEqual(data["mode"], "100% Local Offline")

    def test_ai_chat_endpoint(self):
        payload = {"message": "Xin chào NOVA AI"}
        res = self.client.post("/api/ai/chat", json=payload)
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertIn("reply", data)
        self.assertTrue(len(data["reply"]) > 0)

    def test_ai_sheet_formula_endpoint(self):
        payload = {"query": "Tính tổng cột B từ B1 đến B5", "active_cell": "B6"}
        res = self.client.post("/api/ai/sheet-formula", json=payload)
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertIn("formula", data)
        self.assertTrue(data["formula"].startswith("="))

    def test_ai_slide_generate_endpoint(self):
        payload = {"topic": "Chuyển Đổi Số Doanh Nghiệp", "num_slides": 2}
        res = self.client.post("/api/ai/slide-generate", json=payload)
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertIn("slides", data)
        self.assertTrue(len(data["slides"]) >= 1)

    def test_ai_full_doc_fix_endpoint(self):
        payload = {"html_content": "<h1>tieu de bao cao</h1><p>day la doan van can sua loi chinh ta</p>"}
        res = self.client.post("/api/ai/full-doc-fix", json=payload)
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "ok")
        self.assertIn("fixed_html", data)

    def test_ai_sheet_assist_endpoint(self):
        payload = {"action": "formula", "query": "Tính tổng điểm từ A1 đến A10", "active_cell": "A11"}
        res = self.client.post("/api/ai/sheet-assist", json=payload)
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertIn("formula", data)

if __name__ == "__main__":
    unittest.main()


